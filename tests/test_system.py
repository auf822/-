import io
import os
import tempfile
import unittest
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta, timezone

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from app.main import app
from app.db import init_db, transaction, create_user, now

PASSWORD = 'Testing2026!!'


class SystemTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.original = os.environ.get('DATABASE_PATH')
        os.environ['DATABASE_PATH'] = self.temp.name + '/test.sqlite3'
        init_db()
        self.clients = []
        t = datetime.now(timezone.utc)
        with transaction() as c:
            self.year = c.execute("INSERT INTO years(name,current,stage,r1_start,r1_end,wish_limit) VALUES('测试学年',1,'round1',?,?,2)", ((t-timedelta(days=1)).isoformat(), (t+timedelta(days=1)).isoformat())).lastrowid
            self.ids = {}
            for username, role in [('admin','admin'),('counselor','counselor'),('secretary','secretary'),('t1','advisor'),('t2','advisor'),('t0','advisor'),('s1','student'),('s2','student'),('s3','student')]:
                self.ids[username] = create_user(c, username, PASSWORD, username, role, '计算机', '教授' if role=='advisor' else '', score=90, year_id=self.year if role=='student' else None)
            for key, total in [('t1',1),('t2',2),('t0',0)]:
                c.execute('INSERT INTO quotas VALUES(?,?,?)', (self.year, self.ids[key], total))

    def tearDown(self):
        for client in self.clients:
            client.close()
        if self.original is None:
            os.environ.pop('DATABASE_PATH', None)
        else:
            os.environ['DATABASE_PATH'] = self.original
        self.temp.cleanup()

    def login(self, name):
        client = TestClient(app)
        self.clients.append(client)
        r = client.post('/api/login', json={'username':name,'password':PASSWORD})
        self.assertEqual(r.status_code, 200, r.text)
        client.headers['X-CSRF-Token'] = r.json()['csrf']
        return client

    def apply(self, client, teacher='t1', rank=1):
        return client.post('/api/applications', json={'advisor_id':self.ids[teacher],'rank':rank})

    def accept(self, client, application):
        return client.post(f'/api/applications/{application}/review', json={'decision':'accept','note':''})

    def close_first(self):
        result = self.login('counselor').post('/api/rounds/1/close')
        self.assertEqual(result.status_code, 200, result.text)

    def open_second(self):
        t = datetime.now(timezone.utc)
        result = self.login('counselor').post('/api/rounds/2/open', json={'start':t.isoformat(),'end':(t+timedelta(days=1)).isoformat()})
        self.assertEqual(result.status_code, 200, result.text)

    def workbook(self, headers, rows):
        w = Workbook()
        w.active.append(headers)
        for row in rows:
            w.active.append(row)
        output = io.BytesIO()
        w.save(output)
        return output.getvalue()

    def test_student_payload_hides_quotas_and_zero_advisors(self):
        client = self.login('s1')
        data = client.get('/api/advisors').json()
        self.assertEqual(len(data), 2)
        for item in data:
            self.assertFalse({'total','remaining','matched'} & set(item))
        dashboard = client.get('/api/dashboard').json()
        self.assertFalse({'total','remaining','majors'} & set(dashboard))
        self.assertNotIn('major_caps', dashboard['year'])
        self.assertEqual(self.apply(client,'t0').status_code, 409)
        self.assertEqual(client.get('/api/users').status_code, 403)
        self.assertEqual(client.get('/api/export/results').status_code, 403)

    def test_advisor_sees_only_own_applicants(self):
        a1 = self.apply(self.login('s1'), 't1').json()['id']
        a2 = self.apply(self.login('s2'), 't2').json()['id']
        teacher = self.login('t1')
        self.assertEqual([a['id'] for a in teacher.get('/api/applications').json()], [a1])
        self.assertEqual(self.accept(teacher, a2).status_code, 404)
        self.assertEqual(teacher.get('/api/advisors').status_code, 403)
        self.assertEqual(teacher.get('/api/users').status_code, 403)
        self.assertEqual(teacher.get('/api/export/quotas').status_code, 403)

    def test_rejection_releases_wish_limit(self):
        s = self.login('s1')
        a1 = self.apply(s).json()['id']
        self.assertEqual(self.apply(s, 't2', 2).status_code, 200)
        self.assertEqual(self.apply(s, 't1', 1).status_code, 409)
        self.assertEqual(self.login('t1').post(f'/api/applications/{a1}/review', json={'decision':'reject','note':'研究方向不符'}).status_code, 200)
        self.assertEqual(self.apply(s, 't1', 1).status_code, 200)
        self.assertEqual(s.get('/api/dashboard').json()['pending'], 2)

    def test_acceptance_locks_and_invalidates_other_choices(self):
        s = self.login('s1')
        a = self.apply(s).json()['id']
        b = self.apply(s, 't2', 2).json()['id']
        self.assertEqual(self.accept(self.login('t1'), a).status_code, 200)
        self.assertEqual(self.apply(s,'t2').status_code, 409)
        self.assertEqual(s.delete(f'/api/applications/{a}').status_code, 409)
        self.assertEqual(self.accept(self.login('t2'), b).status_code, 409)
        self.assertEqual(len(s.get('/api/matches').json()), 1)
        self.assertEqual(self.login('t1').get('/api/dashboard').json()['remaining'], 0)
        for name in ['admin','secretary','t1','s1']:
            result = self.login(name).post('/api/matches/adjust', json={'student_id':self.ids['s1'],'advisor_id':None,'reason':'特殊退选'})
            self.assertEqual(result.status_code, 403)

    def test_concurrent_acceptances_do_not_overbook(self):
        a = self.apply(self.login('s1')).json()['id']
        b = self.apply(self.login('s2')).json()['id']
        t1, t2 = self.login('t1'), self.login('t1')
        with ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(lambda pair: self.accept(*pair).status_code, [(t1,a),(t2,b)]))
        self.assertEqual(sorted(results), [200,409])
        self.assertEqual(len(t1.get('/api/matches').json()), 1)

    def test_concurrent_different_teachers_match_student_once(self):
        student = self.login('s1')
        a = self.apply(student,'t1',1).json()['id']
        b = self.apply(student,'t2',2).json()['id']
        t1, t2 = self.login('t1'), self.login('t2')
        with ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(lambda pair: self.accept(*pair).status_code, [(t1,a),(t2,b)]))
        self.assertEqual(sorted(results), [200,409])
        self.assertEqual(len(student.get('/api/matches').json()), 1)

    def test_second_round_manual_and_frozen_eligibility(self):
        s1,s2 = self.login('s1'),self.login('s2')
        a = self.apply(s1).json()['id']
        self.accept(self.login('t1'), a)
        self.close_first()
        self.assertEqual(self.apply(s2,'t2').status_code, 409)
        self.assertEqual(self.login('s2').get('/api/dashboard').json()['year']['stage'], 'round1_closed')
        counselor = self.login('counselor')
        self.assertEqual(counselor.post('/api/matches/adjust',json={'student_id':self.ids['s1'],'advisor_id':None,'reason':'退选申请'}).status_code,200)
        self.open_second()
        self.assertEqual(self.apply(s1,'t2').status_code,403)
        self.assertEqual(self.apply(s2,'t1').status_code,200)
        self.assertEqual(self.login('secretary').put('/api/quotas/'+str(self.ids['t1']),json={'total':3}).status_code,403)
        self.assertEqual(counselor.put('/api/quotas/'+str(self.ids['t1']),json={'total':3}).status_code,200)

    def test_deadlines_block_submit_edit_and_review(self):
        s = self.login('s1')
        a = self.apply(s).json()['id']
        with transaction() as c:
            c.execute('UPDATE years SET r1_end=?', ((datetime.now(timezone.utc)-timedelta(seconds=1)).isoformat(),))
        self.assertEqual(self.apply(s, 't2').status_code,409)
        self.assertEqual(s.put(f'/api/applications/{a}',json={'advisor_id':self.ids['t2'],'rank':1}).status_code,409)
        self.assertEqual(self.accept(self.login('t1'),a).status_code,409)
        self.assertEqual(s.get('/api/applications').json()[0]['status'],'expired')

    def test_quota_lower_bound_and_major_cap_rollback(self):
        secretary = self.login('secretary')
        a = self.apply(self.login('s1')).json()['id']
        self.accept(self.login('t1'),a)
        self.assertEqual(secretary.put(f'/api/quotas/{self.ids["t1"]}',json={'total':0}).status_code,409)
        with transaction() as c:
            c.execute('UPDATE years SET major_caps=?', ('{"计算机":3}',))
        self.assertEqual(secretary.put(f'/api/quotas/{self.ids["t2"]}',json={'total':3}).status_code,409)
        self.assertEqual(next(a for a in secretary.get('/api/advisors').json() if a['id']==self.ids['t2'])['total'],2)

    def test_failed_adjustment_does_not_lose_existing_match(self):
        a = self.apply(self.login('s1')).json()['id']
        self.accept(self.login('t1'),a)
        self.close_first()
        result=self.login('counselor').post('/api/matches/adjust',json={'student_id':self.ids['s1'],'advisor_id':self.ids['t0'],'reason':'特殊调整'})
        self.assertEqual(result.status_code,409)
        self.assertEqual(self.login('s1').get('/api/matches').json()[0]['advisor_id'],self.ids['t1'])

    def test_profile_requires_review_and_only_latest_request(self):
        teacher=self.login('t1')
        old=teacher.post('/api/profile',json={'bio':'旧版本','direction':'旧方向'}).json()['id']
        new=teacher.post('/api/profile',json={'bio':'新版本','direction':'新方向'}).json()['id']
        student=self.login('s1')
        self.assertEqual(student.get('/api/advisors').json()[0]['bio'],'')
        counselor=self.login('counselor')
        self.assertEqual(counselor.post(f'/api/profile-requests/{old}/review',json={'decision':'accept'}).status_code,409)
        self.assertEqual(counselor.post(f'/api/profile-requests/{new}/review',json={'decision':'accept'}).status_code,200)
        self.assertEqual(student.get('/api/advisors').json()[0]['bio'],'新版本')

    def test_csrf_and_origin_enforced(self):
        s=self.login('s1')
        del s.headers['X-CSRF-Token']
        self.assertEqual(self.apply(s).status_code,403)
        self.assertEqual(s.post('/api/login',json={'username':'s1','password':PASSWORD},headers={'Origin':'https://evil.example'}).status_code,403)

    def test_permission_revocation_and_no_cross_role_grants(self):
        counselor=self.login('counselor')
        admin=self.login('admin')
        self.assertEqual(admin.put(f'/api/users/{self.ids["counselor"]}/grants',json={'grants':[]}).status_code,200)
        self.assertEqual(counselor.get('/api/users').status_code,403)
        self.assertEqual(counselor.post('/api/rounds/1/close').status_code,403)
        self.assertEqual(admin.put(f'/api/users/{self.ids["secretary"]}/grants',json={'grants':['matches']}).status_code,400)

    def test_disable_and_password_reset_invalidate_sessions(self):
        student=self.login('s1')
        admin=self.login('admin')
        self.assertEqual(admin.post(f'/api/users/{self.ids["s1"]}/reset-password',json={'password':'Replacement2026!!'}).status_code,200)
        self.assertEqual(student.get('/api/me').status_code,401)
        other=self.login('s2')
        self.assertEqual(admin.put(f'/api/users/{self.ids["s2"]}',json={'name':'s2','major':'计算机','enabled':False}).status_code,200)
        self.assertEqual(other.get('/api/me').status_code,401)

    def test_import_atomic_and_reports_row_numbers(self):
        secretary=self.login('secretary')
        content=self.workbook(['工号','姓名','专业','职称','招生名额'],[['t1','t1','计算机','教授',2],['t2','名字错误','计算机','教授',3]])
        result=secretary.post('/api/import/quotas',content=content)
        self.assertEqual(result.status_code,400)
        self.assertEqual(result.json()['detail']['errors'][0]['row'],3)
        self.assertEqual(secretary.get('/api/advisors').json()[0]['total'],1)
        content=self.workbook(['工号','姓名','专业','职称','招生名额'],[['t1','t1','计算机','教授',2],['t2','t2','计算机','教授',3]])
        self.assertEqual(secretary.post('/api/import/quotas',content=content).json()['count'],2)

    def test_student_import_requires_valid_password_and_duplicate_protection(self):
        counselor=self.login('counselor')
        content=self.workbook(['学号','姓名','专业','成绩','初始密码'],[['new1','新人','计算机',92,PASSWORD],['s1','重复','计算机',91,PASSWORD]])
        self.assertEqual(counselor.post('/api/import/students',content=content).status_code,400)
        self.assertFalse(any(x['username']=='new1' for x in counselor.get('/api/users').json()))
        content=self.workbook(['学号','姓名','专业','成绩','初始密码'],[['new1','新人','计算机',92,PASSWORD]])
        self.assertEqual(counselor.post('/api/import/students',content=content).status_code,200)
        self.assertEqual(self.login('new1').get('/api/me').status_code,200)

    def test_excel_rejects_formulas_and_export_neutralizes_injection(self):
        content=self.workbook(['工号','姓名','专业','职称','招生名额'],[['t1','t1','计算机','教授','=1+2']])
        self.assertEqual(self.login('secretary').post('/api/import/quotas',content=content).status_code,400)
        with transaction() as c:
            c.execute('UPDATE users SET name=? WHERE id=?', ('=HYPERLINK("https://evil.example")', self.ids['t1']))
        response=self.login('admin').get('/api/export/quotas')
        self.assertEqual(response.status_code,200)
        w=load_workbook(io.BytesIO(response.content))
        self.assertEqual(w.active['B2'].data_type,'s')
        self.assertTrue(w.active['B2'].value.startswith("'="))

    def test_archive_and_next_year_are_isolated(self):
        self.close_first()
        self.open_second()
        counselor=self.login('counselor')
        self.assertEqual(counselor.post('/api/rounds/2/close').status_code,200)
        self.assertEqual(counselor.post('/api/results/archive').status_code,409)
        self.assertEqual(counselor.post('/api/results/approve').status_code,200)
        self.assertEqual(counselor.post('/api/results/archive').status_code,200)
        self.assertEqual(counselor.post('/api/matches/adjust',json={'student_id':self.ids['s1'],'advisor_id':self.ids['t1'],'reason':'归档调整'}).status_code,409)
        self.assertEqual(self.login('admin').post('/api/years',json={'name':'下一学年'}).status_code,200)
        self.assertEqual(self.login('admin').get('/api/dashboard').json()['students'],0)
        self.assertEqual(self.login('secretary').get('/api/advisors').json()[0]['total'],0)
        self.assertEqual(len(self.login('admin').get('/api/advisors').json()),3)

    def test_batch_suggestion_confirm_and_atomic_failure(self):
        self.close_first()
        counselor=self.login('counselor')
        plan=counselor.get('/api/matching-plan').json()
        self.assertEqual(len(plan),3)
        self.assertEqual(counselor.get('/api/matches').json(),[])
        failed=counselor.post('/api/matches/batch',json={'assignments':[{'student_id':self.ids['s1'],'advisor_id':self.ids['t1']},{'student_id':self.ids['s2'],'advisor_id':self.ids['t1']}],'reason':'确认建议'})
        self.assertEqual(failed.status_code,409)
        self.assertEqual(counselor.get('/api/matches').json(),[])
        pairs=[{'student_id':p['student_id'],'advisor_id':p['advisor_id']} for p in plan]
        self.assertEqual(counselor.post('/api/matches/batch',json={'assignments':pairs,'reason':'辅导员确认同专业调剂建议'}).json()['count'],3)

    def test_audit_append_only_and_hash_chain(self):
        admin=self.login('admin')
        self.apply(self.login('s1'))
        self.assertTrue(admin.get('/api/audit/verify').json()['valid'])
        self.assertGreaterEqual(len(admin.get('/api/audit').json()),3)
        with self.assertRaises(Exception):
            with transaction() as c:
                c.execute("UPDATE audit SET detail='{}' WHERE id=1")
        with self.assertRaises(Exception):
            with transaction() as c:
                c.execute('DELETE FROM audit')
        self.assertEqual(self.login('secretary').get('/api/audit').status_code,403)

    def test_full_roundtrip_result_export_contains_actual_matches(self):
        s=self.login('s1')
        a=self.apply(s).json()['id']
        self.accept(self.login('t1'),a)
        self.close_first()
        self.open_second()
        s2=self.login('s2')
        b=self.apply(s2,'t2').json()['id']
        self.accept(self.login('t2'),b)
        counselor=self.login('counselor')
        counselor.post('/api/rounds/2/close')
        counselor.post('/api/results/approve')
        r=counselor.get('/api/export/results')
        w=load_workbook(io.BytesIO(r.content),data_only=True)
        self.assertEqual(w['师生配对总表'].max_row,3)
        self.assertEqual(w['导师招生汇总']['F2'].value,0)
        self.assertEqual(w['未匹配学生'].max_row,2)
        self.assertEqual({w['师生配对总表']['A2'].value,w['师生配对总表']['A3'].value},{'s1','s2'})

    def test_archive_snapshot_survives_next_year_profile_changes(self):
        application = self.apply(self.login('s1')).json()['id']
        self.accept(self.login('t1'), application)
        self.close_first()
        self.open_second()
        counselor = self.login('counselor')
        counselor.post('/api/rounds/2/close')
        counselor.post('/api/results/approve')
        counselor.post('/api/results/archive')
        admin = self.login('admin')
        admin.post('/api/years', json={'name':'新学年'})
        self.assertEqual(admin.put(f'/api/users/{self.ids["t1"]}', json={'name':'改名后导师','major':'计算机','title':'教授','enabled':True}).status_code,200)
        self.assertEqual(admin.post(f'/api/years/{self.year}/activate').status_code,200)
        self.assertEqual(admin.get('/api/matches').json()[0]['advisor_name'],'t1')
        workbook = load_workbook(io.BytesIO(admin.get('/api/export/results').content))
        self.assertEqual(workbook['师生配对总表']['E2'].value,'t1')
        self.assertEqual(self.login('t1').get('/api/history').json()[0]['advisor_name'],'t1')

    def test_login_rate_limit_is_persisted_for_failed_requests(self):
        client = TestClient(app)
        self.clients.append(client)
        for _ in range(8):
            self.assertEqual(client.post('/api/login',json={'username':'s1','password':'wrong-password'}).status_code,401)
        self.assertEqual(client.post('/api/login',json={'username':'s1','password':PASSWORD}).status_code,429)

    def test_invalid_time_returns_validation_error(self):
        admin = self.login('admin')
        response = admin.put('/api/rules',json={'wish_limit':3,'start':'not a date','end':'also not a date'})
        self.assertEqual(response.status_code,400)


if __name__ == '__main__':
    unittest.main()
