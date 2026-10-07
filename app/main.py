import hashlib
import io
import json
import os
import secrets
import sqlite3
import threading
import time
import zipfile
from contextlib import asynccontextmanager
from pathlib import Path
from urllib.parse import urlparse
from typing import Literal

from fastapi import FastAPI, Depends, Request, Response, HTTPException
from fastapi.responses import FileResponse, StreamingResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field, ConfigDict
from openpyxl import Workbook, load_workbook

from .db import init_db, transaction, now, audit, notify, password_hash, password_ok, create_user, DEFAULT_GRANTS
from .business import (fail, require, current_year, instant, sync_deadlines, active_round,
                       student_eligible, remaining, check_capacity, change_quota,
                       validate_major_caps, quota_permission, invalidate_approval, match_student)

STATIC = Path(__file__).parent / 'static'
stop = threading.Event()


def deadline_worker():
    while not stop.wait(10):
        with transaction() as c:
            sync_deadlines(c)


@asynccontextmanager
async def lifespan(app):
    init_db()
    stop.clear()
    worker = threading.Thread(target=deadline_worker, daemon=True)
    worker.start()
    yield
    stop.set()
    worker.join(timeout=12)


app = FastAPI(title='研途 · 研究生导师双向选择系统', version='1.0.0', lifespan=lifespan)


@app.middleware('http')
async def security_headers(request, call_next):
    if request.method not in ('GET', 'HEAD', 'OPTIONS'):
        origin = request.headers.get('origin')
        if origin and urlparse(origin).netloc != request.headers.get('host'):
            return JSONResponse({'detail': '请求来源不受信任'}, status_code=403)
    response = await call_next(request)
    response.headers['X-Content-Type-Options'] = 'nosniff'
    response.headers['X-Frame-Options'] = 'DENY'
    response.headers['Referrer-Policy'] = 'same-origin'
    response.headers['Content-Security-Policy'] = "default-src 'self'; style-src 'self'; script-src 'self'; img-src 'self' data:; frame-ancestors 'none'; base-uri 'self'; form-action 'self'"
    if request.url.path.startswith('/api/'):
        response.headers['Cache-Control'] = 'no-store'
    return response


@app.exception_handler(sqlite3.IntegrityError)
async def integrity_error(request, exc):
    return JSONResponse({'detail': '数据冲突：账号或记录已存在，请刷新后重试'}, status_code=409)


def session_hash(token):
    return hashlib.sha256(token.encode()).hexdigest()


def authenticated(request: Request):
    token = request.cookies.get('selection_session', '')
    with transaction() as c:
        sync_deadlines(c)
        s = c.execute('SELECT s.*,u.* FROM sessions s JOIN users u ON u.id=s.user_id WHERE token_hash=? AND expires>? AND enabled=1',
                      (session_hash(token), int(time.time()))).fetchone()
        if not s:
            fail('请先登录，或登录已过期', 401)
        if request.method not in ('GET', 'HEAD') and not secrets.compare_digest(request.headers.get('x-csrf-token', ''), s['csrf']):
            fail('会话校验失败，请刷新页面重试', 403)
        return dict(s)


def public_user(u):
    return {k: u[k] for k in ('id', 'username', 'name', 'role', 'major', 'title', 'score', 'year_id', 'enabled')} | {'grants': json.loads(u['grants'])}


def fresh_user(c, u):
    row = c.execute('SELECT * FROM users WHERE id=? AND enabled=1', (u['id'],)).fetchone()
    if not row:
        fail('账号已停用', 401)
    return row


class Input(BaseModel):
    model_config = ConfigDict(extra='forbid', str_strip_whitespace=True)


class Login(Input):
    username: str = Field(min_length=1, max_length=80)
    password: str = Field(min_length=1, max_length=200)


@app.post('/api/login')
def login(body: Login, request: Request, response: Response):
    key = hashlib.sha256(f'{request.client.host}|{body.username}'.encode()).hexdigest()
    failed = False
    with transaction() as c:
        attempts = c.execute('SELECT * FROM login_attempts WHERE key=?', (key,)).fetchone()
        if attempts and int(time.time()) - attempts['started'] < 900 and attempts['count'] >= 8:
            fail('登录尝试过多，请15分钟后重试', 429)
        u = c.execute('SELECT * FROM users WHERE username=?', (body.username,)).fetchone()
        if not u or not u['enabled'] or not password_ok(body.password, u['password_hash']):
            count = attempts['count'] + 1 if attempts and int(time.time()) - attempts['started'] < 900 else 1
            started = attempts['started'] if count > 1 else int(time.time())
            c.execute('INSERT INTO login_attempts VALUES(?,?,?) ON CONFLICT(key) DO UPDATE SET count=excluded.count,started=excluded.started', (key, count, started))
            audit(c, u['id'] if u else None, 'auth.login_failed', {'username': body.username})
            failed = True
        else:
            c.execute('DELETE FROM login_attempts WHERE key=?', (key,))
            c.execute('DELETE FROM sessions WHERE expires<=?', (int(time.time()),))
            token, csrf = secrets.token_urlsafe(40), secrets.token_urlsafe(32)
            c.execute('INSERT INTO sessions VALUES(?,?,?,?)', (session_hash(token), u['id'], csrf, int(time.time()) + 8 * 3600))
            audit(c, u['id'], 'auth.login', {})
            result = {'user': public_user(u), 'csrf': csrf}
    if failed:
        fail('账号或密码错误', 401)
    response.set_cookie('selection_session', token, httponly=True, samesite='strict', secure=os.environ.get('COOKIE_SECURE') == '1', max_age=28800)
    return result


@app.get('/api/me')
def me(u=Depends(authenticated)):
    return {'user': public_user(u), 'csrf': u['csrf']}


@app.post('/api/logout')
def logout(request: Request, response: Response, u=Depends(authenticated)):
    with transaction() as c:
        c.execute('DELETE FROM sessions WHERE token_hash=?', (u['token_hash'],))
        audit(c, u['id'], 'auth.logout', {})
    response.delete_cookie('selection_session')
    return {'ok': True}


class Password(Input):
    old_password: str = Field(max_length=200)
    new_password: str = Field(min_length=10, max_length=200)


@app.post('/api/password')
def password(body: Password, response: Response, u=Depends(authenticated)):
    with transaction() as c:
        actor = fresh_user(c, u)
        if not password_ok(body.old_password, actor['password_hash']):
            fail('当前密码不正确')
        c.execute('UPDATE users SET password_hash=? WHERE id=?', (password_hash(body.new_password), u['id']))
        c.execute('DELETE FROM sessions WHERE user_id=?', (u['id'],))
        audit(c, u['id'], 'auth.password_change', {})
    response.delete_cookie('selection_session')
    return {'ok': True}


@app.get('/api/years')
def years(u=Depends(authenticated)):
    with transaction() as c:
        rows = list(c.execute('SELECT * FROM years ORDER BY id DESC'))
        if u['role'] in ('student', 'advisor'):
            return [{k: r[k] for k in ('id', 'name', 'stage', 'current')} for r in rows]
        return [dict(r) | {'major_caps': json.loads(r['major_caps'])} for r in rows]


class YearCreate(Input):
    name: str = Field(min_length=1, max_length=50)


@app.post('/api/years')
def new_year(body: YearCreate, u=Depends(authenticated)):
    with transaction() as c:
        require(fresh_user(c, u), ['admin'])
        old = c.execute('SELECT * FROM years WHERE current=1').fetchone()
        if old and old['stage'] != 'archived':
            fail('请先审核并归档当前学年，再创建新学年', 409)
        c.execute('UPDATE years SET current=0')
        uid = c.execute('INSERT INTO years(name,current) VALUES(?,1)', (body.name,)).lastrowid
        audit(c, u['id'], 'year.create', {'id': uid, 'name': body.name})
    return {'id': uid}


@app.post('/api/years/{year_id}/activate')
def activate_year(year_id: int, u=Depends(authenticated)):
    with transaction() as c:
        require(fresh_user(c, u), ['admin'])
        y = c.execute('SELECT * FROM years WHERE id=?', (year_id,)).fetchone()
        old = c.execute('SELECT * FROM years WHERE current=1').fetchone()
        if not y:
            fail('学年不存在', 404)
        if old and old['stage'] not in ('archived', 'preparation'):
            fail('进行中的学年不能切换', 409)
        c.execute('UPDATE years SET current=0')
        c.execute('UPDATE years SET current=1 WHERE id=?', (year_id,))
        audit(c, u['id'], 'year.activate', {'year': year_id})
    return {'ok': True}


class Rules(Input):
    wish_limit: int = Field(ge=1, le=10)
    start: str
    end: str
    major_caps: dict[str, int] = Field(default_factory=dict)


def validate_window(start, end):
    a, b = instant(start), instant(end)
    if a >= b or b.timestamp() <= time.time():
        fail('结束时间必须晚于开始时间和当前时间')
    return a.isoformat(), b.isoformat()


@app.put('/api/rules')
def rules(body: Rules, u=Depends(authenticated)):
    start, end = validate_window(body.start, body.end)
    if any(not k or v < 0 or v > 100000 for k, v in body.major_caps.items()):
        fail('专业名额必须是0至100000的整数')
    with transaction() as c:
        require(fresh_user(c, u), ['admin', 'counselor'], 'rules')
        y = current_year(c)
        if y['stage'] != 'preparation':
            fail('首轮开启后规则不可更改', 409)
        c.execute('UPDATE years SET wish_limit=?,r1_start=?,r1_end=?,major_caps=? WHERE id=?', (body.wish_limit, start, end, json.dumps(body.major_caps), y['id']))
        validate_major_caps(c, current_year(c))
        audit(c, u['id'], 'rules.configure', body.model_dump())
    return {'ok': True}


class Window(Input):
    start: str
    end: str


@app.post('/api/rounds/1/open')
def open_round1(u=Depends(authenticated)):
    with transaction() as c:
        require(fresh_user(c, u), ['counselor'], 'rules')
        y = current_year(c)
        if y['stage'] != 'preparation' or not y['r1_start']:
            fail('请在准备阶段先配置首轮时间', 409)
        validate_window(y['r1_start'], y['r1_end'])
        if not c.execute('SELECT 1 FROM quotas WHERE year_id=? AND total>0', (y['id'],)).fetchone():
            fail('请先由科研秘书配置导师名额', 409)
        c.execute("UPDATE years SET stage='round1' WHERE id=?", (y['id'],))
        people = [r[0] for r in c.execute("SELECT id FROM users WHERE enabled=1 AND (role='advisor' OR (role='student' AND year_id=?))", (y['id'],))]
        notify(c, people, '首轮双选已开启', f'请在配置的时间窗口内完成志愿填报和审核，截止时间：{y["r1_end"]}。')
        audit(c, u['id'], 'round1.open', {'year': y['id']})
    return {'ok': True}


@app.post('/api/rounds/2/open')
def open_round2(body: Window, u=Depends(authenticated)):
    start, end = validate_window(body.start, body.end)
    with transaction() as c:
        require(fresh_user(c, u), ['counselor'], 'rules')
        y = current_year(c)
        if y['stage'] != 'round1_closed':
            fail('首轮结束后才能手动开启二轮，二轮只能开启一次', 409)
        if instant(start) < instant(y['r1_end']):
            fail('二轮开始时间不得早于首轮截止时间')
        c.execute("UPDATE years SET stage='round2',r2_start=?,r2_end=? WHERE id=?", (start, end, y['id']))
        people = [r[0] for r in c.execute('SELECT student_id FROM round2_eligible WHERE year_id=?', (y['id'],))]
        people += [r[0] for r in c.execute("SELECT id FROM users WHERE enabled=1 AND role='advisor'")]
        notify(c, people, '二轮补选已开启', '二轮仅向首轮未匹配学生开放，请在本轮时间窗口内完成操作。')
        audit(c, u['id'], 'round2.open', {'year': y['id'], 'start': start, 'end': end})
    return {'ok': True}


@app.post('/api/rounds/{round_no}/close')
def close_round(round_no: int, u=Depends(authenticated)):
    with transaction() as c:
        require(fresh_user(c, u), ['counselor'], 'rules')
        y = current_year(c)
        if round_no not in (1, 2) or y['stage'] != f'round{round_no}':
            fail('当前轮次不正确', 409)
        c.execute(f'UPDATE years SET r{round_no}_end=? WHERE id=?', (now(), y['id']))
        sync_deadlines(c)
        audit(c, u['id'], 'round.close_manually', {'year': y['id'], 'round': round_no})
    return {'ok': True}


@app.get('/api/advisors')
def advisors(u=Depends(authenticated)):
    with transaction() as c:
        actor = fresh_user(c, u)
        y = current_year(c)
        if actor['role'] == 'advisor':
            fail('导师仅可查看自己的资料和招生数据', 403)
        result = []
        for row in c.execute("SELECT u.id,u.name,u.username,u.major,u.title,u.enabled,p.bio,p.direction,p.projects,p.achievements,p.notice,COALESCE(q.total,0) total FROM users u JOIN profiles p ON p.advisor_id=u.id LEFT JOIN quotas q ON q.advisor_id=u.id AND q.year_id=? WHERE u.role='advisor' AND u.enabled=1 ORDER BY u.id", (y['id'],)):
            item = dict(row)
            if actor['role'] == 'student':
                if item['total'] == 0:
                    continue
                del item['total']
                del item['enabled']
            else:
                item['remaining'] = remaining(c, y['id'], item['id'])
                item['matched'] = item['total'] - item['remaining']
            result.append(item)
        return result


class Quota(Input):
    total: int = Field(ge=0, le=10000)


@app.put('/api/quotas/{advisor_id}')
def quota(advisor_id: int, body: Quota, u=Depends(authenticated)):
    with transaction() as c:
        y = current_year(c)
        quota_permission(fresh_user(c, u), y)
        change_quota(c, y, advisor_id, body.total)
        validate_major_caps(c, y)
        invalidate_approval(c, y)
        audit(c, u['id'], 'quota.update', {'year': y['id'], 'advisor': advisor_id, 'total': body.total})
    return {'ok': True}


class UserCreate(Input):
    username: str = Field(min_length=1, max_length=80, pattern=r'^[A-Za-z0-9_.-]+$')
    password: str = Field(min_length=10, max_length=200)
    name: str = Field(min_length=1, max_length=80)
    role: Literal['admin', 'counselor', 'secretary', 'advisor', 'student']
    major: str = Field(default='', max_length=100)
    title: str = Field(default='', max_length=100)
    score: float = Field(default=0, ge=0, le=1000)


def validate_create(c, actor, body):
    if actor['role'] == 'admin':
        return
    require(actor, ['counselor'], 'people')
    if body.role not in ('student', 'advisor'):
        fail('仅超级管理员可以授权管理岗', 403)


def insert_person(c, body):
    y = current_year(c)
    if y['stage'] == 'archived' and body.role == 'student':
        fail('归档学年不能新增学生', 409)
    if body.role in ('student', 'advisor') and not body.major:
        fail('师生专业不能为空')
    if body.role == 'student' and y['stage'] not in ('preparation', 'round1'):
        fail('首轮截止后不能新增本届学生', 409)
    return create_user(c, body.username, body.password, body.name, body.role, body.major,
                       body.title, body.score, y['id'] if body.role == 'student' else None)


@app.get('/api/users')
def users(role: str = '', u=Depends(authenticated)):
    with transaction() as c:
        actor = fresh_user(c, u)
        require(actor, ['admin', 'counselor'], 'people')
        y = current_year(c)
        rows = c.execute("SELECT * FROM users WHERE (role!='student' OR year_id=?) ORDER BY id", (y['id'],))
        return [public_user(r) for r in rows if (actor['role'] == 'admin' or r['role'] in ('advisor', 'student')) and (not role or r['role'] == role)]


@app.post('/api/users')
def add_user(body: UserCreate, u=Depends(authenticated)):
    with transaction() as c:
        validate_create(c, fresh_user(c, u), body)
        uid = insert_person(c, body)
        audit(c, u['id'], 'user.create', {'id': uid, 'username': body.username, 'role': body.role})
    return {'id': uid}


class UserUpdate(Input):
    name: str = Field(min_length=1, max_length=80)
    major: str = Field(max_length=100)
    title: str = Field(default='', max_length=100)
    score: float = Field(default=0, ge=0, le=1000)
    enabled: bool = True


@app.put('/api/users/{uid}')
def edit_user(uid: int, body: UserUpdate, u=Depends(authenticated)):
    with transaction() as c:
        actor = fresh_user(c, u)
        require(actor, ['admin', 'counselor'], 'people')
        target = c.execute('SELECT * FROM users WHERE id=?', (uid,)).fetchone()
        if not target:
            fail('账号不存在', 404)
        if actor['role'] == 'counselor' and target['role'] not in ('advisor', 'student'):
            fail('只能维护师生账号', 403)
        y = current_year(c)
        if y['stage'] == 'archived' and target['role'] in ('advisor', 'student'):
            fail('归档学年禁止修改师生基础信息', 409)
        if target['role'] == 'student' and target['year_id'] != y['id']:
            fail('不能修改其他学年的学生', 403)
        if target['role'] in ('advisor', 'student') and not body.major:
            fail('专业不能为空')
        if not body.enabled:
            if uid == u['id']:
                fail('不能停用自己的账号')
            if target['role'] == 'admin' and c.execute("SELECT COUNT(*) FROM users WHERE role='admin' AND enabled=1").fetchone()[0] <= 1:
                fail('至少保留一名超级管理员')
            if c.execute('SELECT 1 FROM matches WHERE year_id=? AND (student_id=? OR advisor_id=?)', (y['id'], uid, uid)).fetchone():
                fail('请先由辅导员处理该账号的配对结果', 409)
            c.execute("UPDATE applications SET status='cancelled',reviewed_at=? WHERE year_id=? AND (student_id=? OR advisor_id=?) AND status='pending'", (now(), y['id'], uid, uid))
            c.execute('DELETE FROM sessions WHERE user_id=?', (uid,))
            if target['role'] == 'advisor':
                c.execute('UPDATE quotas SET total=0 WHERE year_id=? AND advisor_id=?', (y['id'], uid))
        c.execute('UPDATE users SET name=?,major=?,title=?,score=?,enabled=? WHERE id=?', (body.name, body.major, body.title, body.score, body.enabled, uid))
        validate_major_caps(c, y)
        invalidate_approval(c, y)
        audit(c, u['id'], 'user.update', {'id': uid, **body.model_dump()})
    return {'ok': True}


class Grants(Input):
    grants: list[str]


@app.put('/api/users/{uid}/grants')
def set_grants(uid: int, body: Grants, u=Depends(authenticated)):
    with transaction() as c:
        require(fresh_user(c, u), ['admin'])
        target = c.execute('SELECT * FROM users WHERE id=?', (uid,)).fetchone()
        if not target or target['role'] not in ('counselor', 'secretary'):
            fail('仅能调整辅导员或科研秘书的专项权限')
        if not set(body.grants).issubset(DEFAULT_GRANTS[target['role']]):
            fail('权限不能突破岗位职责边界')
        c.execute('UPDATE users SET grants=? WHERE id=?', (json.dumps(sorted(set(body.grants))), uid))
        audit(c, u['id'], 'user.grants', {'id': uid, 'grants': body.grants})
    return {'ok': True}


class ResetPassword(Input):
    password: str = Field(min_length=10, max_length=200)


@app.post('/api/users/{uid}/reset-password')
def reset_password(uid: int, body: ResetPassword, u=Depends(authenticated)):
    with transaction() as c:
        require(fresh_user(c, u), ['admin'])
        if not c.execute('SELECT 1 FROM users WHERE id=?', (uid,)).fetchone():
            fail('账号不存在', 404)
        c.execute('UPDATE users SET password_hash=? WHERE id=?', (password_hash(body.password), uid))
        c.execute('DELETE FROM sessions WHERE user_id=?', (uid,))
        audit(c, u['id'], 'user.password_reset', {'id': uid})
    return {'ok': True}


@app.get('/api/applications')
def applications(sort: Literal['time', 'score'] = 'time', u=Depends(authenticated)):
    with transaction() as c:
        actor = fresh_user(c, u)
        require(actor, ['admin', 'counselor', 'advisor', 'student'], 'reports' if actor['role'] == 'counselor' else None)
        y = current_year(c)
        where = 'a.student_id=?' if actor['role'] == 'student' else 'a.advisor_id=?' if actor['role'] == 'advisor' else '1=1'
        args = (y['id'], actor['id']) if actor['role'] in ('student', 'advisor') else (y['id'],)
        order = 's.score DESC,a.created_at ASC' if sort == 'score' else 'a.created_at DESC,a.id DESC'
        rows = c.execute(f'SELECT a.*,s.name student_name,s.username student_no,s.major student_major,s.score,t.name advisor_name FROM applications a JOIN users s ON s.id=a.student_id JOIN users t ON t.id=a.advisor_id WHERE a.year_id=? AND {where} ORDER BY {order}', args)
        return [dict(r) for r in rows]


class Choice(Input):
    advisor_id: int
    rank: int = Field(ge=1, le=10)


@app.post('/api/applications')
def apply(body: Choice, u=Depends(authenticated)):
    with transaction() as c:
        actor = fresh_user(c, u)
        require(actor, ['student'])
        sync_deadlines(c)
        y = current_year(c)
        r = active_round(y)
        student_eligible(c, y, actor, r)
        check_capacity(c, y['id'], body.advisor_id)
        count = c.execute("SELECT COUNT(*) FROM applications WHERE year_id=? AND student_id=? AND status='pending'", (y['id'], actor['id'])).fetchone()[0]
        if count >= y['wish_limit'] or body.rank > y['wish_limit']:
            fail('已达到本学年的志愿填报上限', 409)
        if c.execute("SELECT 1 FROM applications WHERE year_id=? AND student_id=? AND status='pending' AND rank=?", (y['id'], actor['id'], body.rank)).fetchone():
            fail('志愿顺序已被占用，请选择其他顺序', 409)
        uid = c.execute('INSERT INTO applications(year_id,student_id,advisor_id,round,rank,created_at) VALUES(?,?,?,?,?,?)', (y['id'], actor['id'], body.advisor_id, r, body.rank, now())).lastrowid
        notify(c, [body.advisor_id], '收到新的学生意向', f'{actor["name"]}提交了导师意向，请及时审核。')
        audit(c, actor['id'], 'application.submit', {'id': uid, 'advisor': body.advisor_id, 'round': r})
    return {'id': uid}


@app.put('/api/applications/{aid}')
def edit_choice(aid: int, body: Choice, u=Depends(authenticated)):
    with transaction() as c:
        actor = fresh_user(c, u)
        require(actor, ['student'])
        sync_deadlines(c)
        y = current_year(c)
        r = active_round(y)
        student_eligible(c, y, actor, r)
        row = c.execute("SELECT * FROM applications WHERE id=? AND student_id=? AND year_id=? AND status='pending' AND round=?", (aid, actor['id'], y['id'], r)).fetchone()
        if not row:
            fail('志愿不存在或已锁定', 409)
        check_capacity(c, y['id'], body.advisor_id)
        if body.rank > y['wish_limit']:
            fail('志愿顺序超过上限')
        if c.execute("SELECT 1 FROM applications WHERE year_id=? AND student_id=? AND status='pending' AND id!=? AND rank=?", (y['id'], actor['id'], aid, body.rank)).fetchone():
            fail('志愿顺序已被占用', 409)
        c.execute('UPDATE applications SET advisor_id=?,rank=?,created_at=? WHERE id=?', (body.advisor_id, body.rank, now(), aid))
        notify(c, [body.advisor_id], '学生志愿已更新', f'{actor["name"]}更新了志愿，请及时审核。')
        audit(c, actor['id'], 'application.edit', {'id': aid, **body.model_dump()})
    return {'ok': True}


@app.delete('/api/applications/{aid}')
def cancel_choice(aid: int, u=Depends(authenticated)):
    with transaction() as c:
        actor = fresh_user(c, u)
        require(actor, ['student'])
        sync_deadlines(c)
        y = current_year(c)
        r = active_round(y)
        student_eligible(c, y, actor, r)
        row = c.execute("SELECT * FROM applications WHERE id=? AND student_id=? AND year_id=? AND status='pending' AND round=?", (aid, actor['id'], y['id'], r)).fetchone()
        if not row:
            fail('志愿不存在或已锁定', 409)
        c.execute("UPDATE applications SET status='cancelled',reviewed_at=? WHERE id=?", (now(), aid))
        audit(c, actor['id'], 'application.cancel', {'id': aid})
    return {'ok': True}


class Decision(Input):
    decision: Literal['accept', 'reject']
    note: str = Field(default='', max_length=1000)


@app.post('/api/applications/{aid}/review')
def review_application(aid: int, body: Decision, u=Depends(authenticated)):
    with transaction() as c:
        actor = fresh_user(c, u)
        require(actor, ['advisor'])
        sync_deadlines(c)
        y = current_year(c)
        r = active_round(y)
        row = c.execute('SELECT * FROM applications WHERE id=? AND advisor_id=? AND year_id=?', (aid, actor['id'], y['id'])).fetchone()
        if not row:
            fail('志愿不存在', 404)
        if row['status'] != 'pending' or row['round'] != r:
            fail('该志愿已处理或已截止', 409)
        student = c.execute('SELECT * FROM users WHERE id=?', (row['student_id'],)).fetchone()
        student_eligible(c, y, student, r)
        if body.decision == 'accept':
            match_student(c, y, student['id'], actor['id'], r, 'advisor')
            c.execute('UPDATE applications SET note=? WHERE id=?', (body.note, aid))
        else:
            c.execute("UPDATE applications SET status='rejected',reviewed_at=?,note=? WHERE id=?", (now(), body.note, aid))
            notify(c, [student['id']], '导师意向未通过', f'{actor["name"]}未接收本条意向，该志愿名额已释放。{body.note}')
        audit(c, actor['id'], 'application.review', {'id': aid, **body.model_dump()})
    return {'ok': True}


@app.get('/api/matches')
def matches(u=Depends(authenticated)):
    with transaction() as c:
        actor = fresh_user(c, u)
        require(actor, ['admin', 'counselor', 'advisor', 'student'], 'reports' if actor['role'] == 'counselor' else None)
        y = current_year(c)
        where = 'm.student_id=?' if actor['role'] == 'student' else 'm.advisor_id=?' if actor['role'] == 'advisor' else '1=1'
        args = (y['id'], actor['id']) if actor['role'] in ('student', 'advisor') else (y['id'],)
        snapshot = c.execute('SELECT matches FROM year_snapshots WHERE year_id=?', (y['id'],)).fetchone()
        if snapshot:
            key = 'student_id' if actor['role'] == 'student' else 'advisor_id' if actor['role'] == 'advisor' else None
            return [m for m in json.loads(snapshot['matches']) if key is None or m[key] == actor['id']]
        return [dict(r) for r in c.execute(f'SELECT m.*,s.name student_name,s.username student_no,s.major student_major,t.name advisor_name,t.major advisor_major FROM matches m JOIN users s ON s.id=m.student_id JOIN users t ON t.id=m.advisor_id WHERE m.year_id=? AND {where} ORDER BY m.id DESC', args)]


class Adjustment(Input):
    student_id: int
    advisor_id: int | None = None
    reason: str = Field(min_length=2, max_length=1000)


@app.post('/api/matches/adjust')
def adjust(body: Adjustment, u=Depends(authenticated)):
    with transaction() as c:
        require(fresh_user(c, u), ['counselor'], 'matches')
        sync_deadlines(c)
        y = current_year(c)
        if y['stage'] not in ('round1_closed', 'round2', 'completed'):
            fail('人工调剂仅限首轮结束后、学年归档前', 409)
        student = c.execute("SELECT * FROM users WHERE id=? AND role='student' AND year_id=? AND enabled=1", (body.student_id, y['id'])).fetchone()
        if not student:
            fail('学生不属于本届或已停用', 404)
        old = c.execute('SELECT * FROM matches WHERE year_id=? AND student_id=?', (y['id'], student['id'])).fetchone()
        if not old and body.advisor_id is None:
            fail('学生尚未配对，无需解除', 409)
        if old and old['advisor_id'] == body.advisor_id:
            fail('目标导师与当前导师相同')
        c.execute('DELETE FROM matches WHERE year_id=? AND student_id=?', (y['id'], student['id']))
        if body.advisor_id is not None:
            match_student(c, y, student['id'], body.advisor_id, 2 if y['stage'] in ('round2', 'completed') else 1, 'counselor')
        c.execute("UPDATE applications SET status='superseded',reviewed_at=? WHERE year_id=? AND student_id=? AND status='pending'", (now(), y['id'], student['id']))
        invalidate_approval(c, y)
        people = [student['id']] + ([old['advisor_id']] if old else []) + ([body.advisor_id] if body.advisor_id else [])
        notify(c, people, '辅导员已调整配对结果', f'调整原因：{body.reason}')
        audit(c, u['id'], 'match.adjust', {'year': y['id'], 'previous_advisor': old['advisor_id'] if old else None, **body.model_dump()})
    return {'ok': True}


@app.get('/api/matching-suggestions/{sid}')
def suggestions(sid: int, u=Depends(authenticated)):
    with transaction() as c:
        require(fresh_user(c, u), ['counselor'], 'matches')
        y = current_year(c)
        student = c.execute("SELECT * FROM users WHERE id=? AND role='student' AND year_id=?", (sid, y['id'])).fetchone()
        if not student:
            fail('学生不存在', 404)
        rows = []
        for a in c.execute("SELECT u.id,u.name,u.major,p.direction FROM users u JOIN profiles p ON p.advisor_id=u.id WHERE u.role='advisor' AND u.enabled=1"):
            free = remaining(c, y['id'], a['id'])
            if free > 0:
                rows.append(dict(a) | {'remaining': free, 'same_major': a['major'] == student['major']})
        return sorted(rows, key=lambda r: (not r['same_major'], -r['remaining'], r['id']))


@app.get('/api/matching-plan')
def matching_plan(u=Depends(authenticated)):
    """Deterministic suggestions only; explicit counselor confirmation is required."""
    with transaction() as c:
        require(fresh_user(c, u), ['counselor'], 'matches')
        y = current_year(c)
        if y['stage'] not in ('round1_closed', 'round2', 'completed'):
            fail('首轮结束后才能生成调剂建议', 409)
        students = list(c.execute("SELECT * FROM users s WHERE role='student' AND enabled=1 AND year_id=? AND NOT EXISTS(SELECT 1 FROM matches m WHERE m.year_id=? AND m.student_id=s.id) ORDER BY score DESC,id ASC", (y['id'], y['id'])))
        advisors = [dict(r) | {'free': remaining(c, y['id'], r['id'])} for r in c.execute("SELECT id,name,major FROM users WHERE role='advisor' AND enabled=1")]
        plan = []
        for student in students:
            preferences = {r['advisor_id']: r['rank'] for r in c.execute("SELECT advisor_id,MIN(rank) rank FROM applications WHERE year_id=? AND student_id=? AND status IN ('pending','expired') GROUP BY advisor_id", (y['id'], student['id']))}
            candidates = [a for a in advisors if a['free'] > 0 and a['major'] == student['major']]
            if not candidates:
                continue
            candidate = sorted(candidates, key=lambda a: (preferences.get(a['id'], 999), -a['free'], a['id']))[0]
            candidate['free'] -= 1
            plan.append({'student_id': student['id'], 'student_name': student['name'], 'major': student['major'], 'advisor_id': candidate['id'], 'advisor_name': candidate['name'], 'reason': '同专业；优先已有志愿，按成绩与可用名额分配'})
        return plan


class Pair(Input):
    student_id: int
    advisor_id: int


class BatchMatch(Input):
    assignments: list[Pair] = Field(min_length=1, max_length=2000)
    reason: str = Field(min_length=2, max_length=1000)


@app.post('/api/matches/batch')
def batch_match(body: BatchMatch, u=Depends(authenticated)):
    with transaction() as c:
        require(fresh_user(c, u), ['counselor'], 'matches')
        y = current_year(c)
        if y['stage'] not in ('round1_closed', 'round2', 'completed'):
            fail('仅首轮结束后、归档前可以调剂', 409)
        if len({p.student_id for p in body.assignments}) != len(body.assignments):
            fail('学生不能在同一批次重复分配')
        for pair in body.assignments:
            student = c.execute("SELECT * FROM users WHERE id=? AND role='student' AND enabled=1 AND year_id=?", (pair.student_id, y['id'])).fetchone()
            if not student:
                fail('学生不属于本届或已停用', 409)
            match_student(c, y, pair.student_id, pair.advisor_id, 2 if y['stage'] in ('round2', 'completed') else 1, 'counselor')
        audit(c, u['id'], 'match.batch', {'year': y['id'], **body.model_dump()})
    return {'count': len(body.assignments)}


@app.post('/api/results/approve')
def approve_results(u=Depends(authenticated)):
    with transaction() as c:
        require(fresh_user(c, u), ['counselor'], 'reports')
        y = current_year(c)
        if y['stage'] != 'completed':
            fail('请先完成并关闭二轮，再审核最终结果', 409)
        c.execute('UPDATE years SET approved_at=? WHERE id=?', (now(), y['id']))
        audit(c, u['id'], 'results.approve', {'year': y['id']})
    return {'ok': True}


@app.post('/api/results/archive')
def archive_results(u=Depends(authenticated)):
    with transaction() as c:
        require(fresh_user(c, u), ['counselor'], 'reports')
        y = current_year(c)
        if y['stage'] != 'completed' or not y['approved_at']:
            fail('请先审核最终结果，再归档', 409)
        frozen_matches = [dict(r) for r in c.execute('SELECT m.*,s.name student_name,s.username student_no,s.major student_major,t.name advisor_name,t.major advisor_major FROM matches m JOIN users s ON s.id=m.student_id JOIN users t ON t.id=m.advisor_id WHERE m.year_id=? ORDER BY m.id DESC', (y['id'],))]
        c.execute('INSERT INTO year_snapshots(year_id,matches,sheets,stats) VALUES(?,?,?,?)', (y['id'], json.dumps(frozen_matches, ensure_ascii=False), json.dumps(result_sheets(c, y, ''), ensure_ascii=False), json.dumps(management_stats(c, y), ensure_ascii=False)))
        c.execute("UPDATE years SET stage='archived' WHERE id=?", (y['id'],))
        audit(c, u['id'], 'results.archive', {'year': y['id']})
    return {'ok': True}


class Profile(Input):
    bio: str = Field(default='', max_length=5000)
    direction: str = Field(default='', max_length=2000)
    projects: str = Field(default='', max_length=5000)
    achievements: str = Field(default='', max_length=5000)
    notice: str = Field(default='', max_length=2000)


@app.get('/api/profile')
def my_profile(u=Depends(authenticated)):
    require(u, ['advisor'])
    with transaction() as c:
        p = c.execute('SELECT * FROM profiles WHERE advisor_id=?', (u['id'],)).fetchone()
        requests = [dict(r) | {'content': json.loads(r['content'])} for r in c.execute('SELECT * FROM profile_requests WHERE advisor_id=? ORDER BY id DESC', (u['id'],))]
        return {'published': dict(p), 'requests': requests}


@app.post('/api/profile')
def submit_profile(body: Profile, u=Depends(authenticated)):
    with transaction() as c:
        require(fresh_user(c, u), ['advisor'])
        c.execute("UPDATE profile_requests SET status='superseded' WHERE advisor_id=? AND status='pending'", (u['id'],))
        uid = c.execute('INSERT INTO profile_requests(advisor_id,content,created_at) VALUES(?,?,?)', (u['id'], json.dumps(body.model_dump(), ensure_ascii=False), now())).lastrowid
        notify(c, [r[0] for r in c.execute("SELECT id FROM users WHERE role='counselor' AND enabled=1")], '导师资料待审核', f'{u["name"]}提交了公示资料更新。')
        audit(c, u['id'], 'profile.submit', {'id': uid})
    return {'id': uid}


@app.get('/api/profile-requests')
def profile_requests(u=Depends(authenticated)):
    with transaction() as c:
        require(fresh_user(c, u), ['counselor'], 'profiles')
        return [dict(r) | {'content': json.loads(r['content'])} for r in c.execute('SELECT p.*,u.name advisor_name FROM profile_requests p JOIN users u ON u.id=p.advisor_id ORDER BY p.id DESC')]


@app.post('/api/profile-requests/{pid}/review')
def review_profile(pid: int, body: Decision, u=Depends(authenticated)):
    with transaction() as c:
        require(fresh_user(c, u), ['counselor'], 'profiles')
        row = c.execute('SELECT * FROM profile_requests WHERE id=?', (pid,)).fetchone()
        if not row or row['status'] != 'pending':
            fail('资料不存在或已处理', 409)
        if body.decision == 'accept':
            p = Profile(**json.loads(row['content']))
            c.execute('UPDATE profiles SET bio=?,direction=?,projects=?,achievements=?,notice=? WHERE advisor_id=?', (*p.model_dump().values(), row['advisor_id']))
        c.execute('UPDATE profile_requests SET status=?,review_note=? WHERE id=?', ('approved' if body.decision == 'accept' else 'rejected', body.note, pid))
        notify(c, [row['advisor_id']], '个人公示资料审核完成', '资料已发布。' if body.decision == 'accept' else f'资料未通过审核：{body.note}')
        audit(c, u['id'], 'profile.review', {'id': pid, **body.model_dump()})
    return {'ok': True}


@app.get('/api/dashboard')
def dashboard(u=Depends(authenticated)):
    with transaction() as c:
        sync_deadlines(c)
        actor = fresh_user(c, u)
        y = current_year(c)
        result = {'year': dict(y) | {'major_caps': json.loads(y['major_caps'])}}
        if actor['role'] == 'student':
            match = c.execute('SELECT m.*,u.name advisor_name,u.major advisor_major FROM matches m JOIN users u ON u.id=m.advisor_id WHERE m.year_id=? AND m.student_id=?', (y['id'], actor['id'])).fetchone()
            pending = c.execute("SELECT COUNT(*) FROM applications WHERE year_id=? AND student_id=? AND status='pending'", (y['id'], actor['id'])).fetchone()[0]
            eligible = actor['year_id'] == y['id'] and (y['stage'] != 'round2' or bool(c.execute('SELECT 1 FROM round2_eligible WHERE year_id=? AND student_id=?', (y['id'], actor['id'])).fetchone()))
            snapshot = c.execute('SELECT matches FROM year_snapshots WHERE year_id=?', (y['id'],)).fetchone()
            frozen = next((m for m in json.loads(snapshot['matches']) if m['student_id'] == actor['id']), None) if snapshot else None
            result.update(match=frozen or (dict(match) if match else None), pending=pending, eligible=eligible)
            # Students never receive professional caps or quota statistics.
            result['year'].pop('major_caps')
        elif actor['role'] == 'advisor':
            quota = c.execute('SELECT total FROM quotas WHERE year_id=? AND advisor_id=?', (y['id'], actor['id'])).fetchone()
            result.update(total=quota['total'] if quota else 0, remaining=remaining(c, y['id'], actor['id']),
                          pending=c.execute("SELECT COUNT(*) FROM applications WHERE year_id=? AND advisor_id=? AND status='pending'", (y['id'], actor['id'])).fetchone()[0])
        else:
            snapshot = c.execute('SELECT stats FROM year_snapshots WHERE year_id=?', (y['id'],)).fetchone()
            result.update(json.loads(snapshot['stats']) if snapshot else management_stats(c, y))
        return result


def management_stats(c, y):
    return {'students': c.execute("SELECT COUNT(*) FROM users WHERE role='student' AND year_id=? AND enabled=1", (y['id'],)).fetchone()[0],
            'advisors': c.execute("SELECT COUNT(*) FROM users WHERE role='advisor' AND enabled=1").fetchone()[0],
            'matched': c.execute('SELECT COUNT(*) FROM matches WHERE year_id=?', (y['id'],)).fetchone()[0],
            'total': c.execute('SELECT COALESCE(SUM(total),0) FROM quotas WHERE year_id=?', (y['id'],)).fetchone()[0],
            'pending': c.execute("SELECT COUNT(*) FROM applications WHERE year_id=? AND status='pending'", (y['id'],)).fetchone()[0],
            'majors': [dict(r) for r in c.execute("SELECT s.major,COUNT(*) students,SUM(CASE WHEN m.id IS NOT NULL THEN 1 ELSE 0 END) matched FROM users s LEFT JOIN matches m ON m.student_id=s.id AND m.year_id=? WHERE s.role='student' AND s.year_id=? AND s.enabled=1 GROUP BY s.major", (y['id'], y['id']))]}


@app.get('/api/history')
def history(u=Depends(authenticated)):
    require(u, ['advisor', 'student'])
    with transaction() as c:
        where = 'm.advisor_id=?' if u['role'] == 'advisor' else 'm.student_id=?'
        rows = [dict(r) for r in c.execute(f'SELECT y.name year_name,m.*,s.name student_name,s.username student_no,s.major student_major,a.name advisor_name FROM matches m JOIN years y ON y.id=m.year_id JOIN users s ON s.id=m.student_id JOIN users a ON a.id=m.advisor_id WHERE {where} ORDER BY y.id DESC,m.id DESC', (u['id'],))]
        archives = {r['year_id']: json.loads(r['matches']) for r in c.execute('SELECT year_id,matches FROM year_snapshots')}
        for row in rows:
            frozen = next((m for m in archives.get(row['year_id'], []) if m['id'] == row['id']), None)
            if frozen:
                row.update(frozen)
        return rows


@app.get('/api/notifications')
def notifications(u=Depends(authenticated)):
    with transaction() as c:
        return [dict(r) for r in c.execute('SELECT * FROM notifications WHERE user_id=? ORDER BY id DESC LIMIT 100', (u['id'],))]


@app.post('/api/notifications/read')
def mark_notifications(u=Depends(authenticated)):
    with transaction() as c:
        c.execute('UPDATE notifications SET read=1 WHERE user_id=?', (u['id'],))
    return {'ok': True}


class Reminder(Input):
    audience: Literal['students', 'advisors']
    message: str = Field(min_length=1, max_length=1000)


@app.post('/api/reminders')
def reminder(body: Reminder, u=Depends(authenticated)):
    with transaction() as c:
        require(fresh_user(c, u), ['counselor'], 'reminders')
        y = current_year(c)
        if body.audience == 'students':
            people = [r[0] for r in c.execute("SELECT id FROM users s WHERE role='student' AND enabled=1 AND year_id=? AND NOT EXISTS(SELECT 1 FROM matches m WHERE m.student_id=s.id AND m.year_id=?)", (y['id'], y['id']))]
        else:
            people = [r[0] for r in c.execute("SELECT DISTINCT advisor_id FROM applications WHERE year_id=? AND status='pending'", (y['id'],))]
        notify(c, people, '辅导员工作提醒', body.message)
        audit(c, u['id'], 'reminder.send', {'audience': body.audience, 'count': len(people), 'message': body.message})
    return {'count': len(people)}


@app.get('/api/audit')
def logs(u=Depends(authenticated)):
    require(u, ['admin'])
    with transaction() as c:
        return [dict(r) | {'detail': json.loads(r['detail'])} for r in c.execute('SELECT a.*,u.name actor_name FROM audit a LEFT JOIN users u ON u.id=a.actor_id ORDER BY a.id DESC LIMIT 300')]


@app.get('/api/audit/verify')
def verify_audit(u=Depends(authenticated)):
    require(u, ['admin'])
    with transaction() as c:
        previous, count = '0' * 64, 0
        for row in c.execute('SELECT * FROM audit ORDER BY id'):
            digest = hashlib.sha256(f'{previous}|{row["actor_id"]}|{row["action"]}|{row["detail"]}|{row["created_at"]}'.encode()).hexdigest()
            if row['prev_hash'] != previous or row['hash'] != digest:
                return {'valid': False, 'broken_at': row['id']}
            previous, count = digest, count + 1
        return {'valid': True, 'count': count}


def result_sheets(c, y, major=''):
    rows = [tuple(r) for r in c.execute('SELECT s.username,s.name,s.major,t.username,t.name,m.round,m.created_at FROM matches m JOIN users s ON s.id=m.student_id JOIN users t ON t.id=m.advisor_id WHERE m.year_id=? AND (?=\'\' OR s.major=?) ORDER BY s.major,s.username', (y['id'], major, major))]
    sheets = [('师生配对总表', ['学号', '学生姓名', '专业', '导师工号', '导师姓名', '轮次', '配对时间'], rows)]
    for specialty in sorted(set(r[2] for r in rows)):
        clean = ''.join(ch for ch in specialty if ch not in '[]:*?/\\')
        sheets.append((f'专业{len(sheets)}-{clean}'[:31], sheets[0][1], [r for r in rows if r[2] == specialty]))
    quota_rows = []
    for a in c.execute("SELECT u.id,u.username,u.name,u.major,COALESCE(q.total,0) total FROM users u LEFT JOIN quotas q ON q.advisor_id=u.id AND q.year_id=? WHERE role='advisor' ORDER BY u.id", (y['id'],)):
        free = remaining(c, y['id'], a['id'])
        quota_rows.append([a['username'], a['name'], a['major'], a['total'], a['total'] - free, free])
    sheets.append(('导师招生汇总', ['工号', '姓名', '专业', '招生名额', '已录取', '剩余名额'], quota_rows))
    unmatched = [tuple(r) for r in c.execute("SELECT s.username,s.name,s.major,s.score FROM users s WHERE role='student' AND year_id=? AND NOT EXISTS(SELECT 1 FROM matches m WHERE m.year_id=? AND m.student_id=s.id)", (y['id'], y['id']))]
    sheets.append(('未匹配学生', ['学号', '姓名', '专业', '成绩'], unmatched))
    return sheets


def excel_response(sheets, filename):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, headers, rows in sheets:
        sheet = workbook.create_sheet(title[:31])
        sheet.append(headers)
        for row in rows:
            sheet.append(["'" + v if isinstance(v, str) and v.startswith(('=', '+', '-', '@')) else v for v in row])
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        from openpyxl.styles import Font, PatternFill, Alignment
        for cell in sheet[1]:
            cell.fill = PatternFill('solid', fgColor='174D43')
            cell.font = Font(color='FFFFFF', bold=True)
            cell.alignment = Alignment(vertical='center')
        sheet.row_dimensions[1].height = 26
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = min(40, max(16, max(len(str(cell.value or '')) * 1.5 for cell in col) + 2))
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return StreamingResponse(buffer, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': f'attachment; filename="{filename}.xlsx"'})


@app.get('/api/export/{kind}')
def export(kind: Literal['results', 'quotas', 'students-template', 'advisors-template', 'personal'], major: str = '', u=Depends(authenticated)):
    with transaction() as c:
        actor = fresh_user(c, u)
        y = current_year(c)
        sheets = []
        if kind == 'quotas':
            if actor['role'] == 'secretary':
                require(actor, ['secretary'], 'quota_reports')
            else:
                require(actor, ['admin', 'counselor'], 'reports')
            rows = [tuple(r) for r in c.execute("SELECT u.username,u.name,u.major,u.title,COALESCE(q.total,0) FROM users u LEFT JOIN quotas q ON q.advisor_id=u.id AND q.year_id=? WHERE role='advisor' AND enabled=1 ORDER BY u.id", (y['id'],))]
            sheets.append(('导师名额填报', ['工号', '姓名', '专业', '职称', '招生名额'], rows))
        elif kind in ('students-template', 'advisors-template'):
            require(actor, ['admin', 'counselor'], 'people')
            headers = ['学号', '姓名', '专业', '成绩', '初始密码'] if kind == 'students-template' else ['工号', '姓名', '专业', '职称', '初始密码']
            sheets.append(('导入模板', headers, []))
        elif kind == 'personal':
            require(actor, ['advisor'])
            rows = [tuple(r) for r in c.execute('SELECT s.username,s.name,s.major,s.score,m.created_at FROM matches m JOIN users s ON s.id=m.student_id WHERE m.year_id=? AND m.advisor_id=?', (y['id'], actor['id']))]
            sheets.append(('个人录取明细', ['学号', '姓名', '专业', '成绩', '录取时间'], rows))
        else:
            require(actor, ['admin', 'counselor'], 'reports')
            snapshot = c.execute('SELECT sheets FROM year_snapshots WHERE year_id=?', (y['id'],)).fetchone()
            sheets = json.loads(snapshot['sheets']) if snapshot else result_sheets(c, y, major)
            if snapshot and major:
                filtered = [r for r in sheets[0][2] if r[2] == major]
                sheets = [(sheets[0][0], sheets[0][1], filtered), *[s for s in sheets[1:] if not s[0].startswith('专业') or (s[2] and s[2][0][2] == major)]]
        audit(c, actor['id'], 'data.export', {'kind': kind, 'year': y['id'], 'major': major})
        return excel_response(sheets, f'{kind}-{y["id"]}')


async def read_excel(request):
    chunks, size = [], 0
    async for chunk in request.stream():
        size += len(chunk)
        if size > 5 * 1024 * 1024:
            fail('文件不能超过5MB', 413)
        chunks.append(chunk)
    raw = b''.join(chunks)
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            if len(z.infolist()) > 200 or sum(i.file_size for i in z.infolist()) > 20 * 1024 * 1024:
                fail('Excel解压后过大', 413)
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
        if len(book.sheetnames) != 1:
            fail('请使用仅包含一个工作表的标准导入模板')
        if book.active.max_column and book.active.max_column > 20:
            fail('模板包含过多列，请使用标准模板')
        rows = []
        for i, row in enumerate(book.active.iter_rows(values_only=True)):
            if i > 2000:
                fail('单次最多导入2000行')
            if any(isinstance(v, str) and v.startswith('=') for v in row):
                fail(f'第{i+1}行包含公式，请改为数值或文本')
            rows.append(row)
        book.close()
        if not rows:
            fail('文件为空')
        headers = list(rows.pop(0))
        if len(set(headers)) != len(headers):
            fail('模板表头不能重复')
        return headers, rows
    except HTTPException:
        raise
    except Exception:
        fail('无法读取文件，请上传有效的.xlsx模板')


@app.post('/api/import/{kind}')
async def import_excel(kind: Literal['quotas', 'students', 'advisors'], request: Request, u=Depends(authenticated)):
    # Authorize before reading potentially expensive files.
    if kind == 'quotas':
        require(u, ['secretary', 'counselor'])
    else:
        require(u, ['admin', 'counselor'], 'people')
    headers, rows = await read_excel(request)
    expected = ['工号', '姓名', '专业', '职称', '招生名额'] if kind == 'quotas' else ['学号', '姓名', '专业', '成绩', '初始密码'] if kind == 'students' else ['工号', '姓名', '专业', '职称', '初始密码']
    if headers != expected:
        fail('表头与标准模板不一致，请下载并填写系统模板')
    errors, parsed, seen = [], [], set()
    with transaction() as c:
        actor = fresh_user(c, u)
        y = current_year(c)
        if kind == 'quotas':
            quota_permission(actor, y)
        else:
            require(actor, ['admin', 'counselor'], 'people')
        for number, row in enumerate(rows, 2):
            if not any(v is not None and v != '' for v in row):
                continue
            record = dict(zip(headers, row))
            try:
                code = str(record.get(expected[0], '') or '').strip()
                if not code or code in seen:
                    fail('工号/学号为空或在文件中重复')
                seen.add(code)
                if kind == 'quotas':
                    target = c.execute("SELECT * FROM users WHERE username=? AND role='advisor' AND enabled=1", (code,)).fetchone()
                    if not target or any(str(record.get(label, '') or '').strip() != target[field] for label, field in [('姓名', 'name'), ('专业', 'major'), ('职称', 'title')]):
                        fail('导师基础信息与系统不一致')
                    value = record.get('招生名额')
                    if isinstance(value, bool) or not isinstance(value, (int, float)) or int(value) != value:
                        fail('招生名额必须为非负整数')
                    validated = Quota(total=int(value))
                    used = c.execute('SELECT COUNT(*) FROM matches WHERE year_id=? AND advisor_id=?', (y['id'], target['id'])).fetchone()[0]
                    if validated.total < used:
                        fail('名额不能小于已录取人数')
                    parsed.append((target['id'], validated.total))
                else:
                    if c.execute('SELECT 1 FROM users WHERE username=?', (code,)).fetchone():
                        fail('此工号/学号已存在')
                    value = UserCreate(username=code, password=str(record.get('初始密码') or ''), name=str(record.get('姓名') or ''), role='student' if kind == 'students' else 'advisor', major=str(record.get('专业') or ''), title=str(record.get('职称') or ''), score=record.get('成绩', 0) or 0)
                    if not value.major:
                        fail('专业不能为空')
                    parsed.append(value)
            except Exception as exc:
                message = exc.detail if isinstance(exc, HTTPException) else '字段格式错误：请检查必填内容、数字及密码长度（至少10位）'
                errors.append({'row': number, 'message': message})
        if errors:
            fail({'message': '导入失败，未写入任何数据', 'errors': errors[:100]})
        if not parsed:
            fail('模板中没有可导入的数据')
        for value in parsed:
            if kind == 'quotas':
                change_quota(c, y, *value)
            else:
                insert_person(c, value)
        validate_major_caps(c, y)
        invalidate_approval(c, y)
        audit(c, actor['id'], 'data.import', {'kind': kind, 'count': len(parsed), 'year': y['id']})
    return {'count': len(parsed)}


@app.get('/api/health')
def health():
    with transaction() as c:
        c.execute('SELECT 1').fetchone()
    return {'status': 'ok'}


@app.get('/')
def index():
    return FileResponse(STATIC / 'index.html')


app.mount('/static', StaticFiles(directory=STATIC, check_dir=False), name='static')
